# scripts/make_synthetic_docs.py
"""One-off generator for the SYNTHETIC documents in data/raw/ (5,000-8,000 words each).

Dev-only: needs reportlab (not part of the app stack; keep it out of requirements.txt).
Existing files are skipped unless --force; --force touches only the five synthetic
files, never the real Telco CSV. Run from the repo root:
    python -m scripts.make_synthetic_docs
"""
from __future__ import annotations

import argparse
import csv
import re
from itertools import combinations, product
from pathlib import Path
from xml.sax.saxutils import escape

import pdfplumber
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from scripts.synth_docs import fam, item, plans, retention, terms
from scripts.synth_rules import (
    ADDONS, AUTOPAY, BANNER, CALLS_PER_AGENT, CONTRACTS, MAX_WORDS, MIN_WORDS, PLANS,
    SUPPORT_CLOSE, SUPPORT_OPEN, TIERS, Sec, assemble, cents, due, internet_cents,
    make_tickets, phone_dollars, plural, usd, words)
from src.config import DATA_RAW

FORBIDDEN = re.compile(r"\b(roaming|student|senior|military|international|prepaid|mobile)\b", re.I)


def facts() -> list[Sec]:
    """Facts sheet: short core plus computed lookup cards."""
    sla = [["Ticket priority", "First response target", "Resolution target"],
           ["Urgent", "1 hour", "8 hours"], ["High", "4 hours", "24 hours"],
           ["Normal", "1 business day", "3 business days"]]
    core = [
        Sec("Company", ["- Brightline Telecom is a fictional residential telecom provider.",
                        f"- Customer support is open {SUPPORT_OPEN}:00 to {SUPPORT_CLOSE}:00 local time, 7 days a week.",
                        "- Retention offers follow the Retention Policy, and prices follow the Plans and Pricing Guide."]),
        Sec("Support service levels", ["First response targets for Urgent and High tickets count support hours only."], sla),
        Sec("Retention operations", ["- The retention team contacts at-risk customers during support hours only.",
                                     f"- Each retention agent handles at most {CALLS_PER_AGENT} outbound retention calls per day.",
                                     "- The retention manager reviews all Critical tier offers within 1 business day."])]
    bun = []
    for plan, lines, contract, auto in product(PLANS, (1, 2, 3, 4), CONTRACTS, (False, True)):
        net, ph = internet_cents(plan, contract), phone_dollars(lines)
        pay = "paid by automatic payment" if auto else "without automatic payment"
        tail = f", and after the {usd(AUTOPAY)} automatic payment credit it is {cents(net + ph * 100 - AUTOPAY * 100)}" if auto else ""
        bun.append(item(f"Bundle total: {plan}, {lines} phone line{plural(lines)}, {contract.lower()}, {pay}",
                        f"{plan} with {lines} phone line{plural(lines)} on a {contract.lower()} contract, {pay}: internet {cents(net)}, phone {usd(ph)}, total {cents(net + ph * 100)}{tail} per month before add-ons and taxes."))
    pairs = [item(f"Add-on pair: {a} and {b}",
                  f"{a} ({usd(ADDONS[a])}) and {b} ({usd(ADDONS[b])}) together cost {usd(ADDONS[a] + ADDONS[b])} per month, or {usd(12 * (ADDONS[a] + ADDONS[b]))} per year.")
             for a, b in combinations(ADDONS, 2)]
    ceil = [item(f"Discount ceiling: {t} tier, {c.lower()} contract",
                 f"A customer on a {c.lower()} contract in the {t} tier has a combined discount ceiling of {p + cp}% ({cp}% contract plus {p}% retention). On Fiber optic Standard that is {cents(internet_cents('Fiber optic Standard', c, p))} per month before add-ons, phone lines and the automatic payment credit.")
            for (t, _, _, p, _, _), (c, (cp, _, _)) in product(TIERS, CONTRACTS.items())]
    staff = [item(f"Retention capacity: {n} agent{plural(n)}",
                  f"A retention team of {n} agent{plural(n)} can make at most {CALLS_PER_AGENT * n} outbound retention calls per day.")
             for n in range(1, 21)]
    sl = [item(f"Deadline: {pr} ticket opened at {h}:00",
               f"{'An' if pr == 'Urgent' else 'A'} {pr} ticket opened at {h}:00 must have a first response by {due(h, add)}. Support hours count only between {SUPPORT_OPEN}:00 and {SUPPORT_CLOSE}:00.")
          for (pr, add), h in product((("Urgent", 1), ("High", 4)), range(SUPPORT_OPEN, SUPPORT_CLOSE))]
    return assemble(core, [
        fam("Lookup A. Internet and phone bundle totals", "Totals use list prices and the contract discount.", bun),
        fam("Lookup B. Add-on pairs", "Add-on prices are never discounted.", pairs),
        fam("Lookup C. Combined discount ceilings", "The combined ceiling never exceeds the 40 percent cap.", ceil),
        fam("Lookup D. Retention team capacity", "Capacity is agents times the daily call limit.", staff),
        fam("Lookup E. First response deadlines", "Only support hours count toward the target.", sl)])


def build_pdf(path: Path, title: str, secs: list[Sec]) -> None:
    """Ruled tables (so pdfplumber can detect them), page-number footer."""
    st = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=st["BodyText"], fontSize=8.5, leading=10)
    story: list = [Paragraph(escape(title), st["Title"]), Paragraph(escape(BANNER), st["Italic"]), Spacer(1, 10)]
    for s in secs:
        story.append(Paragraph(escape(s.heading), st[f"Heading{s.level}"]))
        story += [Paragraph(escape(p), st["BodyText"]) for p in s.paras]
        if s.table:
            width = (A4[0] - 80) / len(s.table[0])
            grid = Table([[Paragraph(escape(c), cell) for c in row] for row in s.table],
                         colWidths=[width] * len(s.table[0]), repeatRows=1)
            grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                                      ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                                      ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [Spacer(1, 6), grid, Spacer(1, 6)]

    def footer(canvas, doc) -> None:
        canvas.setFont("Helvetica", 8)
        canvas.drawString(40, 25, f"Page {doc.page}")

    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=40, rightMargin=40, topMargin=50,
                      bottomMargin=50).build(story, onFirstPage=footer, onLaterPages=footer)


def build_docx(path: Path, title: str, secs: list[Sec]) -> None:
    """Title style, Heading 1/2, List Bullet, Table Grid."""
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(BANNER)
    for s in secs:
        doc.add_heading(s.heading, s.level)
        for p in s.paras:
            if p.startswith("- "):
                doc.add_paragraph(p[2:], style="List Bullet")
            else:
                doc.add_paragraph(p)
        if s.table:
            grid = doc.add_table(rows=len(s.table), cols=len(s.table[0]))
            grid.style = "Table Grid"
            for r, row in enumerate(s.table):
                for c, val in enumerate(row):
                    grid.cell(r, c).text = val
    doc.save(str(path))


def build_md(path: Path, title: str, secs: list[Sec]) -> None:
    """Markdown with # title, ## sections, ### items and pipe tables."""
    lines = [f"# {title}", "", f"> {BANNER}", ""]
    for s in secs:
        lines += [f"{'#' * (s.level + 1)} {s.heading}", ""]
        for p in s.paras:
            lines += [p, ""]
        if s.table:
            lines += ["| " + " | ".join(s.table[0]) + " |", "|" + " --- |" * len(s.table[0])]
            lines += ["| " + " | ".join(r) + " |" for r in s.table[1:]] + [""]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_csv(path: Path) -> None:
    """Write the synthetic tickets."""
    rows = make_tickets()
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _text_of(path: Path) -> str:
    """Re-extract text from a written file, the way the loaders will."""
    if path.suffix == ".pdf":
        with pdfplumber.open(path) as pdf:
            return "\n".join(p.extract_text() or "" for p in pdf.pages)
    if path.suffix == ".docx":
        d = Document(str(path))
        return "\n".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    if path.suffix == ".csv":
        with path.open(encoding="utf-8", newline="") as fh:
            return "\n".join(" ".join(r) for r in csv.reader(fh))
    return path.read_text(encoding="utf-8")


def verify(paths: list[Path]) -> None:
    """Fail loudly if any file is outside the word range or touches a forbidden topic."""
    for path in paths:
        text = _text_of(path)
        n = words(text)
        print(f"{path.name:28}{n:8d} words")
        if not MIN_WORDS <= n <= MAX_WORDS:
            raise RuntimeError(f"{path.name}: {n} words, expected {MIN_WORDS}-{MAX_WORDS}")
        if hit := FORBIDDEN.search(text):
            raise RuntimeError(f"{path.name}: mentions {hit.group(0)!r}, which must stay a gap")


def main(argv: list[str] | None = None) -> int:
    """Create every missing synthetic raw file, then verify all five."""
    ap = argparse.ArgumentParser(description="Generate synthetic raw documents")
    ap.add_argument("--force", action="store_true", help="overwrite the five synthetic files")
    args = ap.parse_args(argv)
    DATA_RAW.mkdir(parents=True, exist_ok=True)
    jobs = [("facts_sheet.md", lambda p: build_md(p, "Brightline Telecom - Facts Sheet", facts())),
            ("retention_policy.pdf", lambda p: build_pdf(p, "Brightline Telecom - Retention Policy", retention())),
            ("service_terms.pdf", lambda p: build_pdf(p, "Brightline Telecom - Service Terms", terms())),
            ("plans_pricing_guide.docx", lambda p: build_docx(p, "Brightline Telecom - Plans and Pricing Guide", plans())),
            ("support_tickets.csv", write_csv)]
    for name, make in jobs:
        path = DATA_RAW / name
        if path.exists() and not args.force:
            print(f"SKIP (exists, use --force to rebuild): {name}")
            continue
        make(path)
        print(f"wrote {path}")
    verify([DATA_RAW / name for name, _ in jobs])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())