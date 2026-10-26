"""Step 2 smoke tests. Synthetic inputs are built in tmp_path (deterministic)."""
from __future__ import annotations

import pytest
from docx import Document

from src import config
from src.docs.chunker import build_chunks, chunk_section
from src.docs.common import Section, clean_pdf_text, rows_to_markdown
from src.docs.loaders import load_all, load_csv_rows, load_docx, load_markdown
from src.docs.pdf_loader import _Block, _Line, _blocks, _join_pages, _sections, load_pdf
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate


def test_table_split_repeats_header() -> None:
    header = "| Plan | Price | Notes |\n| --- | --- | --- |"
    rows = "\n".join(f"| Plan {i:02d} | ${i + 10}.99 | {'x' * 40} |" for i in range(40))
    pieces = chunk_section(Section(f"{header}\n{rows}", "t.md", "markdown", {}), 500)
    assert len(pieces) > 1
    assert all(piece.startswith("| Plan | Price | Notes |") for piece in pieces)
    joined = "\n".join(pieces)
    assert all(f"| Plan {i:02d} |" in joined for i in range(40))


def test_table_label_stays_with_table() -> None:
    text = "**Churn rate by tenure band**\n\n| Band | Rate |\n| --- | --- |\n| 0-12 | 47.4% |"
    pieces = chunk_section(Section(text, "t.md", "markdown", {}), 900)
    assert len(pieces) == 1
    assert pieces[0].startswith("**Churn rate by tenure band**")
    assert "| 0-12 | 47.4% |" in pieces[0]


def test_cid_placeholders_removed() -> None:
    assert "cid" not in clean_pdf_text("(cid:127) Offer A (cid:5)")
    assert "cid" not in rows_to_markdown([["(cid:127) Plan", "Price"], ["Basic", "$20"]])


def test_text_pieces_respect_budget() -> None:
    text = " ".join(f"Sentence number {i} says something useful about billing." for i in range(60))
    pieces = chunk_section(Section(text, "t.pdf", "pdf", {"page": 1}), 400)
    assert len(pieces) > 1
    assert all(len(piece) <= 400 for piece in pieces)


def test_markdown_loader_tracks_heading_path(tmp_path) -> None:
    md = tmp_path / "doc.md"
    md.write_text("# Policy\nintro\n## Discounts\nMax discount is 20%.\n", encoding="utf-8")
    sections = load_markdown(md)
    assert [s.location for s in sections] == [
        {"section": "Policy"},
        {"section": "Policy > Discounts"},
    ]
    assert "20%" in sections[1].text
    assert chunk_section(sections[1], 900)[0].startswith("[Policy > Discounts]")


def test_csv_rows_one_section_per_row(tmp_path) -> None:
    path = tmp_path / "t.csv"
    path.write_text("id,issue\n1,Late bill\n2,\n", encoding="utf-8")
    sections = load_csv_rows(path)
    assert [s.location for s in sections] == [{"row": 2}, {"row": 3}]
    assert "issue: Late bill" in sections[0].text


def test_docx_loader_keeps_headings_and_tables(tmp_path) -> None:
    doc = Document()
    doc.add_heading("Pricing", level=1)
    doc.add_paragraph("Prices are in USD.")
    table = doc.add_table(rows=2, cols=2)
    for (r, c), value in {(0, 0): "Plan", (0, 1): "Price", (1, 0): "Basic", (1, 1): "$20"}.items():
        table.cell(r, c).text = value
    path = tmp_path / "p.docx"
    doc.save(path)
    sections = load_docx(path)
    assert sections[0].location == {"section": "Pricing"}
    assert sections[1].location == {"section": "Pricing", "table": 1}
    assert "| Basic | $20 |" in sections[1].text


def test_build_chunks_metadata_and_unique_ids() -> None:
    sections = [
        Section("a " * 10, "x.md", "markdown", {}),
        Section("b " * 10, "x.md", "markdown", {"section": "S"}),
    ]
    chunks = build_chunks(sections, 900)
    assert [c["chunk_id"] for c in chunks] == ["x-0000", "x-0001"]
    required = {"chunk_id", "text", "source", "location", "doc_type", "synthetic", "n_chars"}
    assert all(required <= set(c) for c in chunks)
    assert all(isinstance(c["location"], dict) for c in chunks)


RAW_DOCS_PRESENT = all((config.DATA_RAW / name).is_file() for name in config.DOC_FILES)


@pytest.mark.skipif(not RAW_DOCS_PRESENT, reason="raw documents not present")
def test_real_documents_chunk_cleanly() -> None:
    chunks = build_chunks(load_all(), config.CHUNK_MAX_CHARS)
    assert {c["source"] for c in chunks} == set(config.DOC_FILES)
    ids = [c["chunk_id"] for c in chunks]
    assert len(ids) == len(set(ids))
    assert all(c["text"].strip() for c in chunks)
    assert all(isinstance(c["location"], dict) for c in chunks)
    pdf_chunks = [c for c in chunks if c["doc_type"] == "pdf"]
    assert pdf_chunks
    assert all(isinstance(c["location"].get("page"), int) for c in pdf_chunks)


def _line(text: str, top: float, size: float = 10.0, page: int = 1) -> _Line:
    return _Line(text, top, top + size, size, page)


def test_blocks_split_on_gap_and_heading() -> None:
    lines = [_line("1. Billing", 50, 18), _line("Payment is due", 90), _line("15 days after", 102),
             _line("Late fee applies.", 124)]
    blocks = _blocks(lines, body=10.0)
    assert [b.text for b in blocks] == ["1. Billing", "Payment is due 15 days after", "Late fee applies."]
    assert [b.heading for b in blocks] == [True, False, False]


def test_join_pages_rejoins_sentence_but_not_headings() -> None:
    blocks = [_Block("You may offer a", 1, 10, False), _Block("12-month price lock.", 2, 10, False),
              _Block("Done.", 2, 10, False), _Block("Next heading", 3, 18, True),
              _Block("Body without end", 3, 10, False), _Block("Body on page four.", 4, 10, False)]
    joined = _join_pages(blocks)
    assert joined[0].text == "You may offer a 12-month price lock."
    assert joined[2].heading and joined[2].text == "Next heading"
    assert joined[3].text == "Body without end Body on page four."


def test_pdf_sections_follow_heading_path_and_start_page() -> None:
    blocks = [_Block("1. Billing", 1, 18, True), _Block("Payment is due in 15 days.", 1, 10, False),
              _Block("Late fee: day 20", 2, 14, True), _Block("A $5 fee applies.", 2, 10, False)]
    sections = _sections(blocks, "x.pdf")
    assert [s.location for s in sections] == [
        {"section": "1. Billing", "page": 1},
        {"section": "1. Billing > Late fee: day 20", "page": 2}]


def test_pdf_without_headings_falls_back_to_pages() -> None:
    blocks = [_Block("First page text.", 1, 10, False), _Block("Second page text.", 2, 10, False)]
    assert [s.location for s in _sections(blocks, "x.pdf")] == [{"page": 1}, {"page": 2}]


def test_real_pdf_drops_page_numbers_and_joins_wrapped_lines(tmp_path) -> None:
    pytest.importorskip("reportlab")
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    st = getSampleStyleSheet()
    body = " ".join(f"Sentence {i} explains the late fee of $5 for invoices." for i in range(12))
    story: list[Flowable] = [Paragraph("Terms", st["Title"]), Paragraph("1. Billing", st["Heading1"]),
             Paragraph(body, st["BodyText"]), Paragraph("Late fee: day 20", st["Heading2"]),
             Paragraph("The fee is $5.", st["BodyText"])]

    def footer(canvas, doc) -> None:
        canvas.setFont("Helvetica", 8)
        canvas.drawString(40, 25, f"Page {doc.page}")

    path = tmp_path / "t.pdf"
    SimpleDocTemplate(str(path)).build(story, onFirstPage=footer, onLaterPages=footer)
    sections = load_pdf(path)
    assert all("Page 1" not in s.text for s in sections)
    billing = next(s for s in sections if s.location.get("section") == "1. Billing")
    assert "\n" not in billing.text and billing.text.endswith("invoices.")
    assert any(s.location.get("section") == "1. Billing > Late fee: day 20" for s in sections)
    assert all(isinstance(s.location["page"], int) for s in sections)